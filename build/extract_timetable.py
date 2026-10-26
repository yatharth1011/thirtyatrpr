import openpyxl, json, re
from openpyxl.utils import get_column_letter

# Source: official "SEE Time Table for B.Tech. Batch 2026 Semester -1st.xlsx" (Academic Section, 3 Aug 2026)
wb = openpyxl.load_workbook('sources/timetables/SEE Time Table for B.Tech. Batch 2026 Semester -1st.xlsx', data_only=True)
ws = wb['Time Table']
legend = json.load(open('build/venue_legend.json'))

def vnorm(code):
    return re.sub(r'[\s\-()]', '', str(code)).upper()

BRANCH_MAP = {
    'AI':'AIDE','AIB':'AIDE','AIDE':'AIDE',
    'CS':'CSE','CSB':'CSE','CSE':'CSE',
    'EE':'EE','EEB':'EE',
    'EP':'EP','EPB':'EP',
    'ES':'ES','ESB':'ES',
    'DA':'DA','DAB':'DA',
    'CH':'CH','CHB':'CH',
    'CE':'CE','CEB':'CE',
    'ME':'ME','MEB':'ME',
    'MM':'MME','MMB':'MME','MME':'MME',
    'MC':'MCB','MCB':'MCB','MNC':'MCB',
    'ICDT':'ICDT','ICB':'ICDT',
}
BRANCH_TOKEN_RE = re.compile(r'^[A-Za-z]{2,4}$')

VENUE_RE = re.compile(
    r'^(SRC-?\d{3}|S-?\d{3}|S\d{3}|M[1-6]|EE-?(SH|\d)|ME-?(SH|\d)|CY-?(SH|\d)|CS-?(SH|\d)|'
    r'SENATE(\s*HALL)?|AUDI(TORIUM)?(\s*FOYER)?|SRC-?107|L/?P|LAB|AeroD|CYSH|EESH|MESH)$',
    re.I)

# Hand-verified roll/batch tokens (see build/tt_dump.txt manual review). The branch code in each
# tuple must be the BRANCH_MAP-normalized form (e.g. 'CSE', 'AIDE') since that's what student.branch
# holds in the roster the app matches against — NOT the sheet's raw roll-number prefix ('CSB','AIB').
ROLL_TOKENS = {
    '2026CSB1239-1273': [('range','CSE',1239,1273)],
    '2026CSB1274-1288&1693,1694 AND ALL DA BATCH': [('range','CSE',1274,1288), ('nums','CSE',[1693,1694]), ('branch','DA')],
    'CSE 1169-1254': [('range','CSE',1169,1254)],
    '2026CSB1169-1203': [('range','CSE',1169,1203)],
    '2026CSB1204-1238': [('range','CSE',1204,1238)],
    '2026CSB1255-1694 & EP & ICDT BATCHS': [('range','CSE',1255,1694), ('branch','EP'), ('branch','ICDT')],
    '2026MNC1537-1571': [('range','MCB',1537,1571)],
    '2026MNC1572-1581, 1696 & 1697 AND ALL EP STUDENTS': [('range','MCB',1572,1581), ('nums','MCB',[1696,1697]), ('branch','EP')],
    'ME BATCH 2026': [('branch','ME')],
    'ES,ICDT ALL STUDENTS OF BATCH 2026': [('branch','ES'), ('branch','ICDT')],
    '2026AIB1067-1099': [('range','AIDE',1067,1099)],
    '2026MNC1502-1536': [('range','MCB',1502,1536)],
    'WHOLE 2026 BATCH OF MM, CE, CH': [('branch','MME'), ('branch','CE'), ('branch','CH')],
    '2026AIB1001-1033': [('range','AIDE',1001,1033)],
    '2026AIB1034-1066': [('range','AIDE',1034,1066)],
}

def classify_token(tok):
    t = tok.strip()
    if t in ROLL_TOKENS:
        return ('roll', ROLL_TOKENS[t])
    tt = t.upper().replace('&', ' & ')
    m = re.match(r'^D([1-4])\s*&\s*D([1-4])$', tt.replace(' ',''))
    if re.match(r'^D[1-4]$', tt):
        return ('division', [tt])
    m2 = re.match(r'^D([1-4])&D([1-4])$', tt.replace(' ', ''))
    if m2:
        return ('division', [f'D{m2.group(1)}', f'D{m2.group(2)}'])
    if re.match(r'^G(1[0-2]|[1-9])$', tt):
        return ('group', [tt])
    if VENUE_RE.match(t):
        return ('venue', t)
    vk = vnorm(t)
    if vk in legend:
        return ('venue', legend[vk]['code'])
    # branch combo?
    parts = re.split(r'[+&]', t)
    parts = [p.strip() for p in parts if p.strip()]
    if parts and all(BRANCH_TOKEN_RE.match(p) and p.upper() in BRANCH_MAP for p in parts):
        return ('branch', sorted(set(BRANCH_MAP[p.upper()] for p in parts)))
    if t.upper() in ('T','L/P','LAB'):
        return ('note', t)
    if re.match(r'^\d{3}$', t):
        return ('venue', f'S-{t}')  # bare room shorthand seen in sheet (e.g. MA101(001) == S-001)
    return ('unknown', t)

# Cells with no tag in the sheet itself, but confirmed by a separate official email — e.g. the
# Centre of Drone & Autonomous Systems' "Tinkering Lab TL101" email (3 Aug 2026) confirms the
# Monday 6pm/Senate slot is the shared TL101 lecture for the whole Tinkering-track cohort
# (AIDE/CSE/MCB/DA/EP per the TLWP DETAILS sheet split), not just an unrelated upperclassman course.
CONFIRMED_UNTAGGED = {
    'TL101(SENATE)': {'branch': ['AIDE', 'CSE', 'MCB', 'DA', 'EP']},
    # Per "'26 TT.xlsx" (student-shared, cross-checked against '26 official AIDE Division 1
    # schedule): Monday 8-9am is D1's PH101 tutorial. That file doesn't commit to a single room
    # either — column B stacks 6 different venues (S001, EE-SH, EE-1, EE-3, CS-1, CS-SH) at this
    # slot, most likely one per academic group across D1+D2 — so we tag the division/time but
    # deliberately drop the venue rather than assume everyone's tutorial is in S001.
    'PH101(S001)': {'division': ['D1'], 'dropVenue': True},
}
# Same email confirms the Tinkering Lab room itself (the sheet's TL101-LAB cells carry a roll-range
# tag but no venue at all): ground floor, Har Gobind Khorana Block (SAB), Room B004.
CONFIRMED_VENUE_BY_SUBJECT_PREFIX = {
    'TL101-LAB': 'B004, Har Gobind Khorana Block (SAB), ground floor',
}

def parse_cell(text):
    text = str(text)
    parens = re.findall(r'\(([^()]*)\)', text)
    # subject = text up to first '(' or first newline
    head = re.split(r'[\n(]', text, 1)[0].strip()
    tags = {'division': set(), 'group': set(), 'branch': set(), 'roll': [], 'notes': []}
    venue = None
    unknowns = []
    for p in parens:
        kind, val = classify_token(p)
        if kind == 'division':
            tags['division'].update(val)
        elif kind == 'group':
            tags['group'].update(val)
        elif kind == 'branch':
            tags['branch'].update(val)
        elif kind == 'roll':
            tags['roll'].extend(val)
        elif kind == 'venue':
            venue = val
        elif kind == 'note':
            tags['notes'].append(val)
        else:
            unknowns.append(val)
    confirmed = CONFIRMED_UNTAGGED.get(text.strip())
    if confirmed:
        tags['branch'].update(confirmed.get('branch', []))
        tags['division'].update(confirmed.get('division', []))
        tags['group'].update(confirmed.get('group', []))
        if confirmed.get('dropVenue'):
            venue = None

    # trailing plain-text venue after newline (e.g. "...\n\nAuditorium" or "...\nSRC-107")
    plain_segs = [s.strip() for s in text.split('\n') if s.strip()]
    for seg in plain_segs[1:]:
        seg_clean = re.sub(r'\([^()]*\)', '', seg).strip()
        seg_fixed = re.sub(r'^SCR-', 'SRC-', seg_clean, flags=re.I)  # one-off typo in source sheet
        if not seg_clean or venue is not None:
            continue
        if re.match(r'^[A-Za-z][A-Za-z .]*$', seg_clean):
            venue = seg_clean
        elif VENUE_RE.match(seg_fixed):
            venue = seg_fixed.upper()
        elif vnorm(seg_fixed) in legend:
            venue = legend[vnorm(seg_fixed)]['code']
    if venue is None:
        for prefix, v in CONFIRMED_VENUE_BY_SUBJECT_PREFIX.items():
            if head.startswith(prefix):
                venue = v
                break
    return {
        'raw': text.replace('\n', ' / '),
        'subject': head,
        'venue': venue,
        'division': sorted(tags['division']),
        'group': sorted(tags['group']),
        'branch': sorted(tags['branch']),
        'roll': tags['roll'],
        'notes': tags['notes'],
        'unknowns': unknowns,
    }

DAY_BLOCKS = [
    ('Monday', 4, 44),
    ('Tuesday', 45, 77),
    ('Wednesday', 78, 112),
    ('Thursday', 113, 147),
    ('Friday', 148, 178),
]

UNIVERSAL_KEYWORDS = ['NCC/NSO/NSS', 'LIFE SKILL']

result = {}
all_unknowns = {}
for day, start, end in DAY_BLOCKS:
    entries = []
    for c in range(2, 14):  # B..M
        if c == 7:  # G = lunch
            continue
        col = get_column_letter(c)
        time_v = ws.cell(3, c).value
        cat_v = ws.cell(start, c).value
        for r in range(start+1, end+1):
            v = ws.cell(r, c).value
            if v in (None, ''):
                continue
            parsed = parse_cell(v)
            parsed['time'] = str(time_v).strip() if time_v else None
            parsed['category'] = str(cat_v).strip() if cat_v else None
            parsed['col'] = col
            parsed['row'] = r
            is_universal = any(k in parsed['raw'].upper().replace('  ',' ') or k in parsed['raw'] for k in UNIVERSAL_KEYWORDS) or any(k.upper() in parsed['raw'].upper() for k in UNIVERSAL_KEYWORDS)
            parsed['universal'] = bool(is_universal)
            entries.append(parsed)
            if parsed['unknowns']:
                for u in parsed['unknowns']:
                    all_unknowns.setdefault(u, []).append(f'{day} {col}{r}')
    result[day] = entries

with open('build/timetable_parsed.json','w') as f:
    json.dump(result, f, indent=1)

print('UNKNOWN TOKENS:')
for u, locs in all_unknowns.items():
    print(repr(u), locs[:5])

# stats
tot = sum(len(v) for v in result.values())
tagged = sum(1 for v in result.values() for e in v if e['division'] or e['group'] or e['branch'] or e['roll'] or e['universal'])
print(f'total cells: {tot}, confidently-tagged: {tagged}')

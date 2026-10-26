import openpyxl, json, re

# Source: the Academic Section's course-floating sheet, "Final Sheet for TT" in
# sep-new/Final TT 01.09.2026(S).xlsx — every course floated this semester (all batches) with its
# slot(s), enrolment, title and instructor(s). Feeds spectate's Courses tab; the actual meeting
# times/rooms come from the timetable grids (build/extract_institute_tt.py for other batches,
# build/merge.py's timetable for freshers), not from here.
# Columns: B code, C-E slot names, F total students, G title, H instructor(s), I remarks.

wb = openpyxl.load_workbook('sources/updates/sep-new/Final TT 01.09.2026(S).xlsx', data_only=True)
ws = wb['Final Sheet for TT']


def clean(v):
    if v is None:
        return None
    s = str(v).strip()
    return None if not s or s.startswith('#N/A') else s


out = {}
for r in range(2, ws.max_row + 1):
    code = clean(ws.cell(r, 2).value)
    if not code or not re.match(r'^[A-Z]{1,4}\d{2,3}[A-Z]?$', code.upper()):
        continue
    code = code.upper()
    students = ws.cell(r, 6).value
    try:
        students = int(float(students))
    except (TypeError, ValueError):
        students = None
    title = clean(ws.cell(r, 7).value)
    remark = clean(ws.cell(r, 9).value)
    floated = not (remark and re.search(r'not\s+floated|no such course', remark, re.I))
    # The export split titles at commas ("Speak, Think, Pitch…" -> "SPEAK" | " THINK"), pushing the
    # rest into the instructor column; a real instructor cell never starts with a space, a spilled
    # title fragment always does. Such titles are dropped (the app's own course names win anyway).
    raw_instr = ws.cell(r, 8).value
    instructor = clean(raw_instr)
    if isinstance(raw_instr, str) and raw_instr.startswith(' '):
        instructor, title = None, None
    rec = {
        'title': title,
        'instructor': instructor,
        'students': students,
        'slots': [s for s in (clean(ws.cell(r, c).value) for c in (3, 4, 5)) if s],
        'floated': floated,
    }
    if remark:
        rec['remark'] = remark
    out[code] = rec

print('courses:', len(out), '| not floated:', sum(1 for v in out.values() if not v['floated']))
with open('build/course_catalog.json', 'w') as f:
    json.dump(out, f, indent=1)
print('wrote build/course_catalog.json')

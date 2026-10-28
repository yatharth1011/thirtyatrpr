#!/usr/bin/env python3
"""MA101 Quiz 2 (28 Oct 2026) seating plan -> build/exam_seating.json["2026-10-28"].

The workbook is several hall sections stacked vertically (Auditorium first, then "MAIN CAMPUS HALL 3 (M3)",
...), each a grid of repeating 4-column blocks (Seat No. | Course Code | Entry No. | S.No.). A section's hall
is the title row above it, so every entry takes the hall of the nearest title above its row — NOT one hall for
the whole sheet (an earlier version assumed the Auditorium for everyone; wrong). Entries whose seat cell is
blank still get their hall, with no seat.

Only some hall sections carry entries (Auditorium, M3-M6); the others (M2, ME-SH, S-xxx, SRC-107, CY-SH) list
seats with no entries beside them. The last 82 entries sit under no title at all (rows 614+), so their hall is
unknown: they are stored with room 'Venue TBC' rather than guessed.

Run from the portal root; idempotent (only the 2026-10-28 key is rewritten).
"""
import json
import re
import openpyxl

SRC = 'sources/updates/oct-new/MA101 Seating plan_Quiz_2026.xlsx'
DEST = 'build/exam_seating.json'
DATE = '2026-10-28'
ENTRY = re.compile(r'^20\d{2}[A-Z]{2,4}\d{4}$')
SEATLIKE = re.compile(r'^[A-Z]-+\d+$')   # a seat like "C--109" (typo in the sheet) must not read as a title
HALL_CODE = re.compile(r'\((M\d)\)')     # "MAIN CAMPUS HALL 3 (M3)" -> the site's own room code M3
UNTITLED_FROM_ROW = 614                  # entries from here down have no title above them in the sheet

ws = openpyxl.load_workbook(SRC, data_only=True)['Seating plan']

titles = {6: 'Auditorium'}               # row 6: "Seating plan for Auditorium, IIT Ropar"
for r in range(7, ws.max_row + 1):
    a = ws.cell(r, 1).value
    if isinstance(a, str) and not a.startswith(('Seat', 'Date')) and not SEATLIKE.match(a.strip()):
        m = HALL_CODE.search(a)
        titles[r] = m.group(1) if m else a.strip()


def hall_for(row):
    if row >= UNTITLED_FROM_ROW:
        return 'Venue TBC'
    hall = None
    for start in sorted(titles):
        if start <= row:
            hall = titles[start]
    return hall


seats = {}
for row in ws.iter_rows():
    for cell in row:
        v = cell.value
        if isinstance(v, str) and ENTRY.match(v.strip()) and cell.column >= 3:
            seat = ws.cell(cell.row, cell.column - 2).value
            rec = {'room': hall_for(cell.row)}
            if seat:
                rec['seat'] = str(seat).strip()
            seats[v.strip()] = rec

data = json.load(open(DEST))
data[DATE] = seats
json.dump(data, open(DEST, 'w'), indent=2, ensure_ascii=False)
by_hall = {}
for rec in seats.values():
    by_hall[rec['room']] = by_hall.get(rec['room'], 0) + 1
print(f'{DATE}: {len(seats)} entries written to {DEST}:', by_hall)

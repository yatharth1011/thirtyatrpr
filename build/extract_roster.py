import openpyxl, json, re

# Source: official "SEE Time Table for B.Tech. Batch 2026 Semester -1st.xlsx" forwarded by the
# Academic Section (UG) on 3 Aug 2026 — this is the institute's own batch-2026-filtered extract
# of the master timetable + roster, superseding the earlier 21.07.26 consolidated workbook.
wb = openpyxl.load_workbook('sources/timetables/SEE Time Table for B.Tech. Batch 2026 Semester -1st.xlsx', data_only=True)
ws = wb['Divisions Details (D1, D2, D3 a']

roster = {}
for row in ws.iter_rows(min_row=3, values_only=True):
    if row[2] is None:
        continue
    entry = str(row[2]).strip().upper()
    branch = (row[5] or '').strip()
    if not branch:
        continue  # rows with notes (e.g. "restrated with 2026 batch") but no clean branch code
    roster[entry] = {
        'name': row[1].strip(),
        'entry': entry,
        'branch': branch,
        'academicGroup': row[7],   # G1..G12
        'division': row[8].strip() if row[8] else None,  # D1..D4
        'ismpGroup': None,
        'mentor': None,
        'mentorEntry': None,
    }

print('roster size', len(roster))

# ISMP list -> ismp group + mentor
wb2 = openpyxl.load_workbook('sources/timetables/ISMP List of B.Tech 2026.xlsm', data_only=True)
ws2 = wb2['Sheet1']
missing = []
matched = 0
for row in ws2.iter_rows(min_row=2, values_only=True):
    if row[0] is None:
        continue
    entry = str(row[2]).strip().upper()
    grp = row[3]
    ismpGroup = int(grp) if grp is not None else None
    if entry in roster:
        roster[entry]['ismpGroup'] = ismpGroup
        roster[entry]['mentor'] = row[4]
        roster[entry]['mentorEntry'] = row[5]
        matched += 1
    else:
        missing.append(entry)

print('ismp matched', matched, 'ismp-only unmatched', len(missing), missing[:10])

no_ismp = [e for e,v in roster.items() if v['ismpGroup'] is None]
print('roster students w/o ismp group:', len(no_ismp))
print(no_ismp[:20])

# Division stats
from collections import Counter
print('division counts', Counter(v['division'] for v in roster.values()))
print('academic group counts', Counter(v['academicGroup'] for v in roster.values()))
print('ismp group counts', Counter(v['ismpGroup'] for v in roster.values()))
print('branch counts', Counter(v['branch'] for v in roster.values()))

with open('build/roster.json','w') as f:
    json.dump(roster, f, indent=1)

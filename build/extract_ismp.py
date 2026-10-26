import openpyxl, json, re
from openpyxl.utils import get_column_letter

wb = openpyxl.load_workbook('sources/timetables/UPDATED ISMP BTech Schedule 2026.xlsx', data_only=True)
ws = wb['Schedule + Venue(updated)']

header = {}
for c in range(5, 17):
    val = ws.cell(4, c).value
    if val:
        m = re.search(r'\d+', str(val))
        if m:
            header[c] = int(m.group())

merge_map = {}
for m in ws.merged_cells.ranges:
    for r in range(m.min_row, m.max_row+1):
        for c in range(m.min_col, m.max_col+1):
            merge_map[(r,c)] = (m.min_row, m.min_col, m.max_row, m.max_col)

def cell_val(r,c):
    key=(r,c)
    if key in merge_map:
        tr, tc, br, bc = merge_map[key]
        return ws.cell(tr,tc).value, (tr,tc,br,bc)
    return ws.cell(r,c).value, (r,c,r,c)

days = []
cur_date = None
cur_day = None
consumed = set()  # global, across all rows: (min_row,min_col,max_row,max_col)

for r in range(5, 103):
    date_v = ws.cell(r, 2).value
    day_v = ws.cell(r, 3).value
    if date_v not in (None, ''):
        cur_date = str(date_v).strip()
    if day_v not in (None, ''):
        cur_day = str(day_v).strip()
    time_v, _ = cell_val(r, 4)

    events = []
    for c in range(5, 17):
        val, rng = cell_val(r, c)
        if val in (None, ''):
            continue
        if rng in consumed:
            continue
        consumed.add(rng)
        tr, tc, br, bc = rng
        groups = [header[cc] for cc in range(tc, bc+1) if cc in header]
        events.append({'groups': groups, 'text': str(val).strip()})
    if not events:
        continue
    days.append({
        'row': r,
        'date': cur_date,
        'day': cur_day,
        'time': str(time_v).strip() if time_v else None,
        'events': events,
    })

print('total time-rows', len(days))
with open('build/ismp_raw.json','w') as f:
    json.dump(days, f, indent=1)

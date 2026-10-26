import openpyxl, json, re

# Source: the Academic Section's whole-institute grid, "FINAL 21.07.26 TIME TABLE FOR 1st SEMESTER
# OF AY-2026-27.xlsx" (sheet "Final Time Table 21.07.2026") — every batch's courses, not just
# freshers'. Only used for ROOM occupancy (spectate's venue finder): which room is booked when, by
# which course. Freshers' own courses are deliberately dropped here — build/merge.py's hand-curated
# fresher timetable (overrides, cancellations, relocations, real lab durations) is far more accurate
# for those than this July draft grid, and keeping both would double-book their rooms.
#
# Layout: row 3 = period headers in columns B..M (G is the lunch column, left empty); column A
# carries the weekday as a vertically merged block; the first row of each block holds slot names
# (PC-1, HSME, ...), not courses. A cell merged across several columns is one multi-period booking
# (labs), so its duration runs to the end of the last merged period. Cells are free text —
# "CODE(VENUE)", "CODE(G7)(VENUE)", "CODE\nVENUE", "CODE/CODE2(VENUE)" — so the venue is whichever
# parenthesised part or line actually looks like a room.

# The Academic Section keeps this sheet up to date as a Google Sheet (iitrpr login); re-download it
# into oct-new/ (or a newer folder) and point SRC at the latest copy. The original 21 Jul file in
# the portal root is the fallback.
import os
SRC_CANDIDATES = ['sources/updates/oct-new/newsave - FINAL 21.07.26 TIME TABLE FOR 1st SEMESTER OF AY-2026-27.xlsx',
                  'sources/timetables/FINAL 21.07.26 TIME TABLE FOR 1st SEMESTER OF AY-2026-27.xlsx']
SRC = next(p for p in SRC_CANDIDATES if os.path.exists(p))
print('source:', SRC)
wb = openpyxl.load_workbook(SRC, data_only=True)
ws = wb['Final Time Table 21.07.2026']

COL_START = {2: 480, 3: 540, 4: 600, 5: 660, 6: 720, 7: 780, 8: 840, 9: 900, 10: 960, 11: 1020, 12: 1080, 13: 1140}
# Day blocks are found from column A (the weekday labels), since rows shift as the sheet is edited.
# Saturday holds no classes — only a room legend — so it's left out.
WEEKDAY_NAMES = ['Monday', 'Tuesday', 'Wednesday', 'Thursday', 'Friday', 'Saturday']
_starts = [(r, str(ws.cell(r, 1).value).strip()) for r in range(1, ws.max_row + 1)
           if str(ws.cell(r, 1).value or '').strip() in WEEKDAY_NAMES]
DAY_BLOCKS = [(name, r, (_starts[i + 1][0] - 1) if i + 1 < len(_starts) else r)
              for i, (r, name) in enumerate(_starts) if name != 'Saturday']

# Fresher courses — covered (better) by merge.py's own timetable. GE103 is CS100's old code;
# GE107/108/109 are the July draft's placeholder fresher labs (no venue anyway).
FRESHER_CODES = {'AI101', 'MA101', 'PH101', 'CY101', 'CS100', 'GE103', 'TL101', 'WP101', 'HS110', 'HS111',
                 'HS112', 'HS113', 'IK101', 'NO101', 'LS101', 'GE107', 'GE108', 'GE109'}

CODE_RE = re.compile(r'^\s*([A-Z]{1,4}\d{2,3}[A-Z]?(?:\s*/\s*[A-Z]{1,4}\d{2,3}[A-Z]?)*)')
VENUE_RE = re.compile(r'^(SRC|SCR|S|EE|CS|CY|ME)-?(\d{3}|\d|SH|EE\d)$|^SENATE( HALL)?$|^AUDI(TORI?UM)?$|^L/P$', re.I)


def canon_venue(v):
    v = v.strip().upper().replace(' ', '')
    if v in ('AUDI', 'AUDITORIUM', 'AUDITORIM'):
        return 'Auditorium'
    if v.startswith('SENATE'):
        return 'SENATE HALL'
    if v == 'L/P':
        return 'L/P'
    m = re.match(r'^(SRC|SCR|S|EE|CS|CY|ME)-?(.+)$', v)
    pre, rest = m.group(1), m.group(2)
    if pre == 'SCR':
        pre = 'SRC'  # the sheet's own typo for SRC-107
    if pre == 'S' and rest.startswith('EE'):
        return 'EE-' + rest[2:]  # "S-EE3" is EE-3
    return pre + '-' + rest


def venues_in(text):
    parts = re.findall(r'\(([^()]*)\)', text) + text.split('\n')
    out = []
    for p in parts:
        for piece in p.split(','):
            piece = piece.strip()
            if VENUE_RE.match(piece):
                cv = canon_venue(piece)
                if cv not in out:
                    out.append(cv)
    return out


merged_span = {}
for rng in ws.merged_cells.ranges:
    if rng.min_row == rng.max_row and rng.min_col >= 2:
        merged_span[(rng.min_row, rng.min_col)] = rng.max_col

out = {}
skipped = []
for day, r0, r1 in DAY_BLOCKS:
    items = []
    for r in range(r0 + 1, r1 + 1):
        for c in range(2, 14):
            v = ws.cell(r, c).value
            if v is None or not str(v).strip():
                continue
            text = str(v)
            m = CODE_RE.match(text.upper())
            if not m:
                continue
            code = re.sub(r'\s+', '', m.group(1))
            if any(part in FRESHER_CODES for part in code.split('/')):
                continue
            venues = venues_in(text)
            if not venues:
                skipped.append((day, text.strip()))
                continue
            last_col = merged_span.get((r, c), c)
            s = COL_START[c]
            e = COL_START[last_col] + 50
            for venue in venues:
                items.append({'code': code, 'venue': venue, 's': s, 'e': e, 'raw': text.strip().replace('\n', ' ')})
    # A course listed twice in the same slot+room (the sheet repeats a few) is one booking.
    seen = set()
    uniq = []
    for it in items:
        k = (it['code'], it['venue'], it['s'])
        if k in seen:
            continue
        seen.add(k)
        uniq.append(it)
    out[day] = sorted(uniq, key=lambda x: (x['s'], x['venue']))

total = sum(len(v) for v in out.values())
print('institute bookings with a room:', total, {d: len(v) for d, v in out.items()})
print('skipped (no room in the cell):', len(skipped))
with open('build/institute_tt.json', 'w') as f:
    json.dump(out, f, indent=1)
print('wrote build/institute_tt.json')

import openpyxl, json, re

# Source: official per-student "HSS_IKS COURSE LIST.xlsx" (Academic Section) — gives each
# student's assigned HS/IKS elective directly, replacing the need to ask the student to pick.
wb = openpyxl.load_workbook('sources/updates/newdata/HSS_IKS COURSE LIST.xlsx', data_only=True)
ws = wb['Sheet1']

assigned = {}
conflicts = set()
for r in range(2, ws.max_row + 1):
    entry = ws.cell(r, 2).value
    course = ws.cell(r, 3).value
    if not entry or not course:
        continue
    entry = re.sub(r'\s+', '', str(entry)).upper()
    course = str(course).strip()
    if entry in assigned and assigned[entry] != course:
        conflicts.add(entry)
        continue
    assigned[entry] = course

for e in conflicts:
    assigned.pop(e, None)

with open('build/hss_iks_assignment.json', 'w') as f:
    json.dump(assigned, f, indent=1)

print('assigned', len(assigned), 'conflicts dropped', len(conflicts))

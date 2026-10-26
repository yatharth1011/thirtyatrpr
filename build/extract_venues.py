import openpyxl, json, re

# Source: official "SEE Time Table for B.Tech. Batch 2026 Semester -1st.xlsx" (Academic Section, 3 Aug 2026)
wb = openpyxl.load_workbook('sources/timetables/SEE Time Table for B.Tech. Batch 2026 Semester -1st.xlsx', data_only=True)
ws = wb['Time Table']

legend = {}
def norm(code):
    return re.sub(r'[\s\-()]', '', str(code)).upper()

for r in range(183, 216):
    newname = ws.cell(r,2).value
    oldname = ws.cell(r,3).value
    desc = ws.cell(r,4).value
    loc = ws.cell(r,6).value
    desc = (desc or '').strip()
    if loc:
        desc = f"{loc.strip()} — {desc}" if desc else loc.strip()
    if newname:
        legend[norm(newname)] = {'code': str(newname).strip(), 'desc': desc, 'alias': str(oldname).strip() if oldname else None}
    if oldname:
        key = norm(oldname)
        if key not in legend:
            legend[key] = {'code': str(oldname).strip(), 'desc': desc, 'alias': str(newname).strip() if newname else None}

# a few extra known venues not in legend table but used in cells
extra = {
    'AUDITORIUM': 'Swami Vivekanand Auditorium', 'AUDI': 'Swami Vivekanand Auditorium',
    'SENATE': 'Senate Hall — M. Visvesvaraya Block, West Wing (280 capacity)',
    'SENATEHALL': 'Senate Hall — M. Visvesvaraya Block, West Wing (280 capacity)',
    'MESH': 'Mechanical Engineering Seminar Hall',
    'CYSH': 'Chemistry Seminar Hall',
    'CSSH': 'CSE Seminar Hall (90 capacity)',
    'EESH': 'EE Seminar Hall — JC Bose Block (80 capacity)',
    'MESH2': 'ME Seminar Hall',
}
for k,v in extra.items():
    if norm(k) not in legend:
        legend[norm(k)] = {'code': k, 'desc': v, 'alias': None}
    else:
        legend[norm(k)]['desc'] = v

# fill in M1-M6 aliases with their SRC-counterpart's (now location-enriched) description
for k, v in legend.items():
    if v['code'] in ('M1','M2','M3','M4','M5','M6') and v['alias']:
        src_entry = legend.get(norm(v['alias']))
        if src_entry:
            v['desc'] = src_entry['desc']

print('legend entries:', len(legend))
for k,v in list(legend.items())[:20]:
    print(k, v)

with open('build/venue_legend.json','w') as f:
    json.dump(legend, f, indent=1)
